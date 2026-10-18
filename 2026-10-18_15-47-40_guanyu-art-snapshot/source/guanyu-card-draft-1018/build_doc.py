from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path('./guanyu-card-draft-1018')
REF=Path('./user-images-1018/image(20261018-044056).png')
FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BOLD='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
# A neutral information-layout diagram, not character art.
S=2
im=Image.new('RGB',(450*S,660*S),'white'); d=ImageDraw.Draw(im)
def rect(box,fill=None,outline='#4f5358',width=2):
    d.rectangle(tuple(int(x*S) for x in box),fill=fill,outline=outline,width=width*S)
def text(x,y,t,size=16,color='#202327',bold=False):
    d.text((x*S,y*S),t,font=ImageFont.truetype(BOLD if bold else FONT,size*S),fill=color)
def center(x,y,t,size=16,color='#202327',bold=False):
    f=ImageFont.truetype(BOLD if bold else FONT,size*S)
    d.text((x*S,y*S),t,font=f,fill=color,anchor='mm')
rect((22,20,428,635),'#eef0f1','#30353a',5)
rect((36,34,414,92),'#ffffff')
text(50,46,'关羽',27,bold=True); text(134,57,'称号待议',16,color='#666666')
rect((336,44,400,81),'#ebedef'); center(368,62,'★ 1',19,bold=True)
rect((37,104,413,491),'#f8f8f8','#8a8f94',1)
rect((48,126,307,480),'#e1e4e6','#7f858a',2)
center(177,267,'人物立绘区域',24,bold=True)
center(177,309,'姿态见人物说明',17)
center(177,339,'局部可压左侧外框',16)
center(177,369,'面部与长须要清楚',16)
rect((286,364,401,417),'#ffffff','#30353a',2)
center(343,383,'被动',17,bold=True); center(343,404,'名称待议',13)
rect((286,426,401,479),'#ffffff','#30353a',2)
center(343,444,'主动',17,bold=True); center(343,466,'名称待议',13)
rect((37,503,413,558),'#ffffff')
center(99,523,'生命 H',17); center(226,523,'攻击 A',17); center(348,523,'防御 D',17)
center(225,546,'字段与数值待议',12,color='#777777')
rect((37,571,413,620),'#f6f6f6')
text(53,584,'等级  L / 上限待议',16)
text(299,585,'黑铁质感',15)
text(39,642,'布局建议  ·  非立绘成品',13,color='#666666')
im.save(OUT/'layout_wireframe.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.73); sec.right_margin=Inches(.73)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.25)
for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','List Bullet','List Number']:
    st=doc.styles[nm]; st.font.name='Noto Sans CJK SC'; st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
normal=doc.styles['Normal']; normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.2; normal.paragraph_format.space_after=Pt(6)
for nm,sz in [('Title',24),('Heading 1',16),('Heading 2',12)]:
    st=doc.styles[nm]; st.font.size=Pt(sz); st.font.bold=True
    st.paragraph_format.space_before=Pt(12 if nm!='Title' else 0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
for nm in ['List Bullet','List Number']:
    st=doc.styles[nm]; st.font.size=Pt(11)
    st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.2
    st.paragraph_format.left_indent=Inches(.19); st.paragraph_format.first_line_indent=Inches(-.16)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('555555')

def p(txt='',style=None):
    return doc.add_paragraph(txt,style)
NEXT_PAGE=False
def h(txt):
    global NEXT_PAGE
    para=p(txt,'Heading 1')
    if NEXT_PAGE: para.paragraph_format.page_break_before=True; NEXT_PAGE=False
    return para
def sh(txt): return p(txt,'Heading 2')
def bullet(txt): return p(txt,'List Bullet')
def page():
    global NEXT_PAGE
    NEXT_PAGE=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for j,x in enumerate(headers): t.rows[0].cells[j].text=x
    for row in rows:
        for j,x in enumerate(row): t.add_row().cells[j].text=x if False else ''
    # rows are added once below; remove placeholder rows above.
    while len(t.rows)>1: t._tbl.remove(t.rows[-1]._tr)
    for row in rows:
        cells=t.add_row().cells
        for j,x in enumerate(row): cells[j].text=x
    for ri,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            tcPr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); mar.append(el)
            tcPr.append(mar)
            if ri==0:
                fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'E6E8EA'); tcPr.append(fill)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1.12
                if j==0: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs: run.font.size=Pt(10); run.bold=(ri==0)
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in t.rows:
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
    p('').paragraph_format.space_after=Pt(0)
    return t

# Page 1
p('一星关羽卡牌设计讨论稿','Title')
p('v0.1  |  2026年10月18日', 'Subtitle')
p('先把一星关羽的形象、卡面和技能方向讨论清楚，再决定具体数值。本稿只覆盖收集养成所需的卡牌设计，不涉及战场实现。')
p('标记说明  【已定】已确认方向  【建议】本稿提案  【待议】尚未决定').runs[0].font.size=Pt(9.5)
h('一星的核心定位')
p('【已定】避祸逃亡、卖豆谋生的关羽，尚未成名，也未正式从军；有乡勇般的气质，隐隐透出侠义。')
bullet('【已定】从涿郡卖绿豆、初遇刘备的故事线起步，一星没有标志性武器。逃亡卖豆是本项目采用的故事设定，不作为正史陈述。')
bullet('【已定】一星为黑铁阶。整体采用非写实的三国幻想厚涂风格；人物、被动和主动都要体现成长。')
bullet('【已定】以你喜欢的金卡总体布局作为参考。低星的材质、装饰和特效需要克制，信息必须清楚。')
h('人物形象与背景')
p('【已定】脸沿用你认可的金卡关羽的辨识度，保留长须、丹凤眼和熟悉的面部特征。年龄观感可讨论，不换成无关的少年脸。')
p('【建议】布衣与日常束装为主，背景取涿郡街市一角，以豆筐提示卖豆身份。金甲、赤兔和高阶神光暂不放进一星。')
p('【已定】天空仅有第二眼才辨出的蛰伏青龙虚影，暗示潜质；不抢人物，不画实体龙，不参与一星技能，也不抢后续成长空间。\n【已定】姿态先用站豆摊旁，一手扶扁担，微侧头看街上。\n【待议】服色、年龄观感、头巾和挽袖等细节继续讨论。')
sh('故事短文草案')
p('【建议】避祸流落涿郡，关羽在市集中卖绿豆谋生。他尚未从军，一身布衣仍掩不住乡勇般的气质，言行间隐隐透出侠义。与刘备初遇时，他的人生才刚刚翻开新的一页。')
p('【待议】卡面称号可暂叫“市井藏锋”，也可以只写“关羽”。这段故事的语气、细节与称号一起确认。')

# Page 2
page(); h('卡面布局与信息层级')
p('【建议】保留人物占主要面积、右侧信息前置的方向。人物可局部压左框；姓名、星级、技能和属性保持完整可读。')
pp=p(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp.add_run().add_picture(str(OUT/'layout_wireframe.png'),height=Inches(4.12))
pp.add_run('    ')
pp.add_run().add_picture(str(REF),height=Inches(4.12))
for inline in doc.inline_shapes:
    dp=inline._inline.docPr
    if not dp.get('descr'): dp.set('descr','左侧为一星卡面信息线框，右侧为用户提供的金色关羽卡牌参考。')
p('左  一星布局线框建议，仅示意区域与层级。  右  用户提供的金卡缩略参考。','Caption')
table(['层级','卡面字段','当前处理'],[
    ['第一眼','关羽姓名  人物  一颗星','【建议】姓名顶部，星级右上，人物为主体。'],
    ['第二眼','被动  主动  基础属性','【建议】右侧放两条技能名，底部放三项属性；字段待议。'],
    ['补充','等级  称号  技能说明','【建议】等级底部；完整说明放详情页，避免挤在卡面。'],
],[.67,2.03,4.1])
p('【待议】阵营、兵种、总战力是否需要，暂不加入。一星黑铁框的纹样、明暗和厚度，等人物形象确认后再细化。')

# Page 3
page(); h('基础数值与技能初稿')
p('【已定】每次升星，基础属性暂按 ×2 设计。本稿只给一星的数值位置；基础绝对值尚未确定。')
table(['拟用字段','一星基础值','口径'],[
    ['生命','H  待定','【建议】正数基础属性；与等级成长分开记录。'],
    ['攻击','A  待定','【建议】作为一星技能的主要关联属性。'],
    ['防御','D  待定','【建议】是否使用此字段仍需确认。'],
],[1.0,1.35,4.45])
p('【待议】生命、攻击、防御只是拟用字段。升星 ×2 与升级增量如何计算、哪些属性参与翻倍，都要再确认；暴击率等百分比属性不直接套用 ×2。')
sh('被动技能  名称待议')
p('【已定】被动最多一个，效果温和。\n【建议】以一个简单的常驻效果表现乡勇底子。“膂力”“胆气”都只是命名候选。')
bullet('触发：获得或解锁该被动后常驻生效；解锁时点待议。')
bullet('效果：攻击小幅提高 p%。p 为待定参数，需按“温和”定位再定强度。')
bullet('计算：建议将“基础攻击”和“被动加成”分开显示，避免升星翻倍时重复计算；百分比是否随培养变化待议。')
p('【待议】被动名称及具体效果尚未确定；“刚直”先作为人物气质，不另设第二个被动。')
sh('主动技能  名称待议')
p('【已定】主动技能一个，符合逃亡卖豆的阶段。\n【建议】短时间提升自身攻击，表现朴素的奋起之力。“奋勇”“奋起”都是命名候选。')
bullet('触发：主动发动，具体发动规则待议。')
bullet('效果：自身攻击提高 q%，持续 T；q、T 和持续单位均待议。')
bullet('限制：消耗、冷却、重复发动和加成叠加规则待议；本稿不规定战场操作方式。')
p('【建议】先确认这一对技能的方向，再逐项补名称和参数。')

# Page 4
page(); h('获取与培养的讨论边界')
p('【建议】一星承担认识关羽和开始培养的作用。先保证获得、查看和成长的关系清楚，再补材料与节奏。')
sh('获取方式')
p('【建议】可考虑在初遇刘备的故事节点固定获得一星关羽，让人物来历与入手时刻对应。')
p('【待议】是否固定获得、能否重复获取，以及重复卡如何处理。此处没有确定抽卡、卡池或付费方式。')
sh('培养与升星')
p('【建议】等级提升负责日常小幅成长，升星负责阶段性变化。人物形象、被动和主动的成长应能被玩家感知。')
p('【待议】一星等级上限、升级消耗、升星条件与材料、已有等级及技能培养是否继承。现在先保留空位，不规定完整养成系统。')
p('【已定】升星基础属性暂按 ×2。哪些加成属于基础值、哪些属于额外效果，需要统一后再填数。')
h('接下来先讨论这两件事')
p('1  先细化形象。以豆摊旁扶扁担、微侧头看街上的姿态为基础，再逐项讨论服色和年龄观感，保留金卡关羽的辨识度。')
p('2  再确认技能。一个温和被动暂用“小幅提升攻击”，一个主动暂用“临时提升攻击”，这两种效果是否符合你心中的一星关羽？名称随后再定。')
p('卡名、具体数值和培养消耗可以随后逐项定。这一版中带【建议】和【待议】的内容，都可以直接修改。')

# Footer page field for a four-page discussion draft.
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run(); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
# Remove inherited decorative paragraph borders from all styles and text.
for el in doc.styles.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
for el in doc.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
doc.core_properties.title='一星关羽卡牌设计讨论稿 v0.1'
doc.core_properties.subject='一星关羽的形象、卡面和技能讨论'
doc.core_properties.author=''
doc.core_properties.keywords='关羽,一星,卡牌,讨论稿'
path=OUT/'一星关羽卡牌设计讨论稿_v0.1.docx'; doc.save(path)
print(path)
